from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


def _norm(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip().lower()


def _rows(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        out.append({headers[i]: row[i] for i in range(len(headers))})
    return out


def reconcile_training(training_path: str, requirements_path: str, output_path: str) -> dict[str, Any]:
    training = _rows(Path(training_path))
    reqs = _rows(Path(requirements_path))
    by_emp: dict[str, dict[str, str]] = {}
    roles: dict[str, str] = {}
    for rec in training:
        emp = str(rec.get("Employee") or "").strip()
        skill = _norm(rec.get("Skill"))
        status = str(rec.get("Status") or "").strip()
        role = str(rec.get("Role") or "").strip()
        if not emp or not skill:
            continue
        by_emp.setdefault(emp, {})[skill] = status
        if role:
            roles[emp] = role
    required_by_role: dict[str, list[str]] = {}
    for rec in reqs:
        role = str(rec.get("Role") or "").strip()
        skill = str(rec.get("RequiredSkill") or "").strip()
        if role and skill:
            required_by_role.setdefault(role, []).append(skill)
    complete = {"certified", "complete", "yes", "trained", "current"}
    gaps = []
    compliant = 0
    for emp, skills in by_emp.items():
        role = roles.get(emp, "")
        for req in required_by_role.get(role, []):
            have = skills.get(_norm(req))
            if have and _norm(have) in complete:
                compliant += 1
            else:
                gaps.append({"employee": emp, "role": role, "skill": req, "status": have or "NOT FOUND", "severity": "CRITICAL" if not have else "HIGH"})
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Gaps"
    ws.append(["Employee", "Role", "Skill", "Status", "Severity"])
    for g in gaps:
        ws.append([g["employee"], g["role"], g["skill"], g["status"], g["severity"]])
    wb.save(out)
    required_total = compliant + len(gaps)
    return {"employees": len(by_emp), "required_pairs": required_total, "compliant": compliant, "gaps": gaps, "gap_count": len(gaps), "match_rate_pct": round(100 * compliant / required_total, 1) if required_total else 100.0, "report_path": str(out)}
